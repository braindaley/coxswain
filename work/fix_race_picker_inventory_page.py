from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt


SPEC_PATH = "/Users/brian/Documents/ChatGPT/Coxswain/outputs/Coxswain_Product_Design_Spec.docx"
OVERVIEW_PATH = Path("/Users/brian/Documents/ChatGPT/Coxswain/work/screen-inventory-overview.png")


doc = Document(SPEC_PATH)

# The final inventory grew by one row when the Race Reference Picker was added.
# Tighten the image just enough to keep its caption on the same landscape page.
inventory = doc.inline_shapes[-1]
inventory_blip = inventory._inline.graphic.graphicData.pic.blipFill.blip
doc.part.related_parts[inventory_blip.embed]._blob = OVERVIEW_PATH.read_bytes()
aspect_ratio = inventory.height / inventory.width
inventory.width = Inches(7.65)
inventory.height = int(inventory.width * aspect_ratio)

for paragraph in doc.paragraphs:
    if paragraph.text.startswith("Figure 33. Complete Coxswain screen inventory:"):
        paragraph.text = (
            "Figure 33. Complete Coxswain screen inventory: 30 current screens and "
            "interaction states, plus 11 planned placeholders."
        )
        paragraph.paragraph_format.space_before = Pt(2)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.keep_together = True
        for run in paragraph.runs:
            run.font.size = Pt(7)
        break

doc.save(SPEC_PATH)
