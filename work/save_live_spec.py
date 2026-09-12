from docx import Document
from pathlib import Path
import shutil
p=Path('outputs/Coxswain_Product_Design_Spec.docx')
shutil.copy2(p,'work/Coxswain_spec_v1_2_backup.docx')
d=Document(p)
for para in d.paragraphs:
 if 'Android companion app' in para.text:
  for run in para.runs: run.text=run.text.replace('v1.2','v1.3')
for t in d.tables:
 for row in t.rows:
  c=row.cells
  if len(c)>=3 and c[1].text in ['Live Row: Free','Duration Setup','Distance Setup']: c[2].text='LOCKED'
  if len(c)>=2 and c[0].text=='1.1' and c[1].text=='Current / locked':c[1].text='Previous / locked'
d.add_page_break()
d.add_heading('08 Live Row Free',level=1)
d.add_paragraph('Approved design added in v1.3. Live Row Free is an open-ended rowing display whose metric selection reflects each user’s needs. All selected numbers share the largest common font size that fits cleanly. The screen has no fixed primary metric.')
sections={
'Display layout':['Use the approved dark immersive background #042C3D and equal-sized metric areas. No hero metric, decorative chart, or bottom navigation.','Use Roboto with tabular numerals. Every selected metric uses the same numeric font size. Calculate the largest shared size that fits the available width and height of every slot, retaining balanced margins and room for labels and units. Do not cap live numbers at the earlier 48–72sp range. Recalculate when metric selection, count, orientation, or available space changes.','Keep the font size stable during normal value changes; reserve width for the expected numeric format. Longer values must remain fully visible.','Keep labels and units compact and readable in #CAD4E1. Retain the 8dp spacing grid, 16dp phone inset, and minimum 48dp touch targets.'],
'Metric configuration':['Edit display opens metric configuration. The approved interactive mockup also allows selecting a metric area to replace its metric. Preserve access to configuration during a session.','Users can replace, add, and remove metric slots, retaining at least one. Preserve saved configurations and separate ordinary-rowing and pace-boat configurations as supported by the existing app.','Available choices follow the current picker: distance, duration, strokes, energy, speed, heart rate, stroke rate, stroke ratio, power, clock time, split, average split, distance delta, and time delta.','The display supports portrait and landscape. Exact automatic grid breakpoints remain implementation details to validate on phones and tablets; the preview’s arrangements are illustrative.'],
'Session controls':['Header: Free Row, a labeled connection indicator using #22C55E for the connected dot, and Edit display.','Pause and End session remain visible side by side below the metric grid. Pause uses primary blue #0B63F6 with white text. End session uses #DCEBFF with #10213F text; finishing a session is distinct from discarding it. Both use the established pill shape.','Pause changes to Resume when paused. Resume returns to the active session.','End session opens a confirmation with Keep rowing and End session. If already paused, the cancel action is Stay paused. Cancel preserves the previous session state. Confirm finishes the session and proceeds to workout review. The full Workout Complete screen remains to be designed.','The preview uses sample values and local interactions; it is not connected to rowing hardware and does not save workout data.'],
'Approval scope':['Approved: configurable equal-area metrics, a shared maximized number size, dark live-row surface, coordinated brand styling, Edit display, Pause or Resume, and End session with confirmation.','Not yet specified: full connection-loss recovery, missing-sensor behavior, detailed Workout Complete design, and device-specific layout validation. These are follow-up states, not implied approvals.'],
'Source and design reference':['Code reviewed: https://github.com/svenmeier/coxswain — WorkoutActivity.java, BindingView.java, BindingDialogFragment.java, and ValueBinding.java.','An editable copy of the approved interactive design is stored beside this document as Coxswain_Live_Row_Free.html. Its orientation and layout controls are design-preview controls, not added product settings.']}
for title,paras in sections.items():
 d.add_heading(title,level=2)
 for x in paras:d.add_paragraph(x)
d.add_heading('Version update',level=2)
d.add_paragraph('v1.3 — Added the approved Live Row Free design and controls. Corrected screen inventory statuses for Duration Setup and Distance Setup to match their existing locked sections. Home, Duration, and Distance designs remain unchanged.')
d.save(p)
shutil.copy2('/Users/brian/.codex/visualizations/2026/09/07/01a07dce-ee45-77f1-a36d-c2efbed2db4d/coxswain-live-display.html','outputs/Coxswain_Live_Row_Free.html')
