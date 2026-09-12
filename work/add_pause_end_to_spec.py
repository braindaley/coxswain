from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Inches, Pt, RGBColor


ROOT = Path('/Users/brian/Documents/ChatGPT/Coxswain')
SPEC = ROOT / 'outputs/Coxswain_Product_Design_Spec.docx'
OVERVIEW = ROOT / 'work/screen-inventory-overview.png'
STATES = [
    (
        '16.31 Pause / End — Paused',
        ROOT / 'work/screen-atlas/31_pause_end_paused.png',
        'Current v2.0 state: Paused workout',
        [
            'Pause freezes every metric, elapsed counter, target, interval position, and race position while preserving the active session in memory.',
            'The frozen metric grid keeps the same order, labels, and equal numeric sizing used by the live screen that opened it.',
            'Resume returns to that exact Live Row mode and state; End Session opens the save confirmation.',
        ],
    ),
    (
        '16.32 Pause / End — End Confirmation',
        ROOT / 'work/screen-atlas/32_pause_end_confirm.png',
        'Current v2.0 state: End and save confirmation',
        [
            'The bottom sheet keeps the paused workout visible behind a scrim and summarizes the completed distance and duration.',
            'End and Save records the partial or completed workout in History and proceeds to Workout Complete.',
            'Stay Paused closes the sheet. Discard Workout is visually secondary and opens a separate destructive confirmation.',
        ],
    ),
    (
        '16.33 Pause / End — Discard Confirmation',
        ROOT / 'work/screen-atlas/33_pause_end_discard.png',
        'Current v2.0 state: Discard confirmation',
        [
            'Discard is never combined with End and Save in a single primary choice.',
            'The second confirmation states that the workout will not appear in History and that the action cannot be undone.',
            'Go Back returns to the End confirmation without resuming or changing session data.',
        ],
    ),
]

doc = Document(SPEC)

inside_pause = False
status_updated = False
register_replacements = {
    'Purpose: Safely pause, resume, finish, or discard an active workout.':
        'Purpose: Freeze an active workout safely, return to it, save its completed portion, or deliberately discard it.',
    'Primary content: Paused state, elapsed status, resume action, end choices.':
        'Primary content: The frozen live metric layout, a clear Paused state, Resume and End Session actions, an end summary, and separate save and discard confirmations.',
    'Primary actions: Resume; End and Save; Discard Workout; cancel.':
        'Primary actions: Resume; End Session; End and Save; Stay Paused; Discard Workout; Go Back.',
    'States and rules: Discard requires explicit confirmation. End and Save proceeds to Workout Complete.':
        'States and rules: Pause freezes all metrics and program progress. Resume restores the exact Free, Target, Interval, Rest, or Race state that launched the screen. End and Save writes the workout to History and proceeds to Workout Complete. Discard requires a second destructive confirmation and returns Home without creating History.',
}

for paragraph in doc.paragraphs:
    if paragraph.text == '12 Pause End':
        inside_pause = True
        continue
    if paragraph.text == '13 Workout Complete':
        inside_pause = False
    if inside_pause and paragraph.text == 'PLANNED':
        paragraph.text = 'LOCKED'
        status_updated = True
    if inside_pause and paragraph.text in register_replacements:
        paragraph.text = register_replacements[paragraph.text]

if not status_updated:
    raise RuntimeError('Pause / End status was not updated')

text_replacements = {
    'Recommended next screens to define and lock: 1) Race Reference Picker, 2) Pause / End, 3) Workout Complete, 4) History and Workout Details, 5) Connect Rower.':
        'Recommended next screens to define and lock: 1) Workout Complete, 2) History and Workout Details, 3) Connect Rower, 4) Settings and Display Customization.',
    'Designed screens: Home, Quick Duration, Quick Distance, Quick Intervals, Segment Entry, Program Builder, My Programs, Workout Library, Program Detail / Start, Race Reference Picker, Live Row Free, Live Row Target, and Live Row Race are represented in the atlas.':
        'Designed screens: Home, Quick Duration, Quick Distance, Quick Intervals, Segment Entry, Program Builder, My Programs, Workout Library, Program Detail / Start, Race Reference Picker, Live Row Free, Live Row Target, Live Row Race, and Pause / End are represented in the atlas.',
    'v2.0 - Consolidated all designed screens into an individual-state visual atlas, documented component changes by interaction, unified quick workout and program setup, simplified interval segments to Duration / Distance / Rest, and incorporated Race Your Best across setup, reference selection, and Live Row.':
        'v2.0 - Consolidated all designed screens into an individual-state visual atlas, documented component changes by interaction, unified quick workout and program setup, simplified interval segments to Duration / Distance / Rest, incorporated Race Your Best across setup, reference selection, and Live Row, and locked the Pause / End safeguards.',
}
found = set()
for paragraph in doc.paragraphs:
    if paragraph.text in text_replacements:
        original = paragraph.text
        paragraph.text = text_replacements[original]
        found.add(original)
if found != set(text_replacements):
    raise RuntimeError('Expected specification summary text was not found')

# Refresh and tighten the final all-screen inventory.
inventory = doc.inline_shapes[-1]
inventory_blip = inventory._inline.graphic.graphicData.pic.blipFill.blip
doc.part.related_parts[inventory_blip.embed]._blob = OVERVIEW.read_bytes()
inventory.width = Inches(7.65)
inventory.height = Inches(6.45)
for paragraph in doc.paragraphs:
    if paragraph.text.startswith('Figure 33. Complete Coxswain screen inventory:'):
        paragraph.text = 'Figure 36. Complete Coxswain screen inventory: 33 current screens and interaction states, plus 10 planned placeholders.'
        paragraph.paragraph_format.space_before = Pt(2)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.keep_together = True
        for run in paragraph.runs:
            run.font.size = Pt(7)
        break

anchor = next(p for p in doc.paragraphs if p.text == '17 v2.0 Consolidated Decisions')
inserted = []
for heading_text, image_path, caption_text, bullets in STATES:
    page_break = doc.add_paragraph()
    page_break.add_run().add_break(WD_BREAK.PAGE)
    inserted.append(page_break)
    inserted.append(doc.add_paragraph(heading_text, style='Heading 2'))
    figure = doc.add_paragraph()
    figure.alignment = WD_ALIGN_PARAGRAPH.CENTER
    figure.add_run().add_picture(str(image_path), width=Inches(3.0))
    inserted.append(figure)
    caption = doc.add_paragraph(caption_text)
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in caption.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(83, 100, 124)
    inserted.append(caption)
    inserted.append(doc.add_paragraph('Interaction changes', style='Heading 3'))
    for text in bullets:
        inserted.append(doc.add_paragraph(text, style='List Bullet'))

for paragraph in inserted:
    anchor._p.addprevious(paragraph._p)

doc.save(SPEC)
print('added Pause / End states to the specification')
