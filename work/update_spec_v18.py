from pathlib import Path
from docx import Document

SPEC = Path('/Users/brian/Documents/ChatGPT/Coxswain/outputs/Coxswain_Product_Design_Spec.docx')
doc = Document(SPEC)

def replace_text(paragraph, old, new):
    if old not in paragraph.text:
        return
    updated = paragraph.text.replace(old, new)
    if paragraph.runs:
        paragraph.runs[0].text = updated
        for run in paragraph.runs[1:]:
            run.text = ''
    else:
        paragraph.add_run(updated)

for paragraph in doc.paragraphs:
    replace_text(paragraph, 'v1.7', 'v1.8')
for section in doc.sections:
    for paragraph in section.header.paragraphs:
        replace_text(paragraph, 'v1.7', 'v1.8')
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                replace_text(paragraph, 'v1.7', 'v1.8')

current_screen = None
for paragraph in doc.paragraphs:
    if paragraph.text.startswith('03 Duration Setup'):
        current_screen = 'quick'
    elif paragraph.text.startswith('04 Distance Setup'):
        current_screen = 'quick'
    elif paragraph.text.startswith('06 Program Builder'):
        current_screen = 'builder'
    elif paragraph.style.name.startswith('Heading 2') and current_screen:
        current_screen = None
    elif current_screen == 'quick' and paragraph.text.startswith('Primary actions:'):
        if 'Save as Program' not in paragraph.text:
            paragraph.text += ' Save as Program stores the current configuration without starting it.'
    elif current_screen == 'builder' and paragraph.text.startswith('Primary content:'):
        paragraph.text = 'Primary content: The shared Workout Setup shell with program name, type choice, configuration, optional Set Goal component, interval segment list when applicable, and live summary.'
    elif current_screen == 'builder' and paragraph.text.startswith('Primary actions:'):
        paragraph.text = 'Primary actions: Save Program; add, edit, reorder, or remove segments; preview summary; cancel.'

if not any(p.text == '15 Unified Workout Setup Experience' for p in doc.paragraphs):
    doc.add_page_break()
    doc.add_heading('15 Unified Workout Setup Experience', level=1)
    doc.add_paragraph('Duration, Distance, Interval Quick Start, and Program Builder use one shared Workout Setup shell. The target selector, rowing presets, Set Goal tiles, goal adjustment panel, and live summary keep the same layout and behavior across every entry point.')
    doc.add_heading('Quick workout context', level=2)
    doc.add_paragraph('Home Quick Start opens the shell with the selected type already active. The title reads Duration Workout, Distance Workout, or Interval Workout. Program Name is hidden. Start Workout is the primary action and creates an active workout. Save as Program is a secondary action that preserves the current configuration for reuse.')
    doc.add_heading('Program context', level=2)
    doc.add_paragraph('Create Program opens the same shell with Program Name visible and the Duration, Distance, and Intervals type selector available. Save Program is the primary action. Editing an owned program uses the same screen with saved values loaded and Save Changes as the primary action.')
    doc.add_heading('Terminology and behavior', level=2)
    doc.add_paragraph('The interface uses Save as Program because a workout is an active or completed session and a program is its reusable definition. Saving never starts a workout. Starting a quick workout does not require saving. Saving from Quick Start transfers the current configuration into Program context so the user can name and confirm it without re-entering values.')
    doc.add_heading('Shared component contract', level=2)
    doc.add_paragraph('Changing between Duration, Distance, and Intervals replaces only the configuration area. Set Goal and Summary remain in stable positions. Each type preserves its draft configuration while the screen is open. Back asks about discarding changes only after the user has modified the setup.')
    doc.add_heading('Version update', level=2)
    doc.add_paragraph('v1.8 - Unified quick Duration, Distance, and Interval setup with Program Builder. Added Save as Program to quick setup while preserving Start Workout as the primary action and the established Program versus Workout terminology.')

doc.save(SPEC)
