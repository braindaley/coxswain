from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches


SPEC = Path('/Users/brian/Documents/ChatGPT/Coxswain/outputs/Coxswain_Product_Design_Spec.docx')
IMAGE = Path('/Users/brian/Documents/ChatGPT/Coxswain/work/screen-atlas/20_program_detail.png')

doc = Document(SPEC)

replacements = {
    'PLANNED': 'LOCKED',
    'Primary content: Name, description, full configuration, ordered segments, goal, estimated totals when meaningful.':
        'Primary content: Program source, name, description, read-only workout configuration, performance goal, completion history, and Race Your Best when compatible history exists.',
    'Primary actions: Start Workout; Edit for owned programs; save copy for library programs; back.':
        'Primary actions: Start Workout or Start Race; view program History; open owned-program actions; back. Library programs use the same review structure without owned-program management actions.',
    'States and rules: Do not present an exact total duration when distance-based segments make it unknowable.':
        'States and rules: Duration, distance, and interval programs reuse the same sections with type-specific values. Do not present an exact duration when distance-based work makes it unknowable. Race Your Best appears only with compatible completed workouts and replaces the ordinary start action while enabled.',
    'Designed screens: Home, Quick Duration, Quick Distance, Quick Intervals, Segment Entry, Program Builder, My Programs, Workout Library, Live Row Free, and Live Row Race are represented in the atlas.':
        'Designed screens: Home, Quick Duration, Quick Distance, Quick Intervals, Segment Entry, Program Builder, My Programs, Workout Library, Program Detail / Start, Live Row Free, and Live Row Race are represented in the atlas.',
    'Figure 23. Complete Coxswain screen inventory: 19 current screens and interaction states, plus 14 planned placeholders.':
        'Figure 24. Complete Coxswain screen inventory: 20 current screens and interaction states, plus 13 planned placeholders.',
}

found = set()
inside_program_detail = False
for paragraph in doc.paragraphs:
    text = paragraph.text.strip()
    if text == '08 Program Detail Start':
        inside_program_detail = True
        continue
    if inside_program_detail and text == 'PLANNED':
        paragraph.text = 'LOCKED'
        found.add('PLANNED')
        inside_program_detail = False
        continue
    if text in replacements and text != 'PLANNED':
        paragraph.text = replacements[text]
        found.add(text)

required = set(replacements)
missing = required - found
if missing:
    raise RuntimeError('Expected text was not found: ' + repr(sorted(missing)))

if not any(p.text.strip() == '16.20 Program Detail / Start' for p in doc.paragraphs):
    anchor = next(p for p in doc.paragraphs if p.text.strip() == '17 v2.0 Consolidated Decisions')
    heading = anchor.insert_paragraph_before('16.20 Program Detail / Start', style='Heading 2')
    heading.paragraph_format.page_break_before = True

    picture = anchor.insert_paragraph_before('', style='Normal')
    picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture.add_run().add_picture(str(IMAGE), width=Inches(3))

    caption = anchor.insert_paragraph_before('', style='Normal')
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = caption.add_run('Current v2.0 state: Program Detail / Start')
    run.italic = True

    anchor.insert_paragraph_before('Interaction changes', style='Heading 3')
    anchor.insert_paragraph_before(
        'The program definition is read-only and adapts to duration, distance, or interval content.',
        style='List Bullet',
    )
    anchor.insert_paragraph_before(
        'Compatible program History enables Race Your Best; the primary action changes from Start Workout to Start Race.',
        style='List Bullet',
    )
    anchor.insert_paragraph_before('', style='Normal')

doc.save(SPEC)
print('added Program Detail / Start to the specification and visual atlas')
