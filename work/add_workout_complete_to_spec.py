from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Inches, Pt, RGBColor


ROOT = Path('/Users/brian/Documents/ChatGPT/Coxswain')
SPEC = ROOT / 'outputs/Coxswain_Product_Design_Spec.docx'
OVERVIEW = ROOT / 'work/screen-inventory-overview.png'
STATES = [
    ('16.34 Workout Complete — Distance', '34_workout_complete_distance.png', 'Distance result', [
        'A fixed-distance workout leads with completion time; the target distance remains explicit in the program label.',
        'New Best compares completion time only with workouts using the identical distance and program definition.',
        'The summary keeps average split, stroke rate, power, and energy compact beneath the result.',
    ]),
    ('16.35 Workout Complete — Duration', '35_workout_complete_duration.png', 'Duration result', [
        'A fixed-duration workout leads with distance achieved during the target time.',
        'New Best compares distance with completed workouts using the identical duration and program definition.',
        'Time does not appear as the best-result value because it is fixed by the program.',
    ]),
    ('16.36 Workout Complete — Intervals', '36_workout_complete_intervals.png', 'Interval result', [
        'The primary result follows the interval ranking rule and excludes Rest segments from performance metrics.',
        'The complete Row and Rest group appears on one line, followed by individual Row results in the interactive screen.',
        'The same structure supports any free-form ordered mix of Duration, Distance, and Rest segments.',
    ]),
    ('16.37 Workout Complete — Free Row', '37_workout_complete_free.png', 'Free Row result', [
        'Free Row leads with distance and states that the open session was ended by the user.',
        'The screen avoids target-completion or early-end language because Free Row has no fixed target.',
        'The completed workout is saved to History and can be opened or repeated.',
    ]),
    ('16.38 Workout Complete — Race', '38_workout_complete_race.png', 'Race result', [
        'Race completion adds a clear outcome and retains the date and result of the fixed saved reference.',
        'Finish margin and time difference make the result understandable without relying on color.',
        'View Details opens the full Race Comparison Details path through the completed workout.',
    ]),
    ('16.39 Workout Complete — Ended Early', '39_workout_complete_ended.png', 'Manually ended target workout', [
        'A manually ended target workout states Ended Early and leads with actual work completed.',
        'The summary shows elapsed duration and percent of target completed without implying the target was achieved.',
        'The partial workout remains saved accurately in History after End and Save.',
    ]),
]

doc = Document(SPEC)

inside_complete = False
status_updated = False
register_replacements = {
    'Purpose: Confirm completion and summarize results.':
        'Purpose: Confirm a completed or manually ended workout and present the result that matters for its workout type.',
    'Primary content: Key measured totals, achieved target, interval results, date/time.':
        'Primary content: Type-specific primary result, program identity, completion status, best or race comparison when applicable, key averages, interval results, and date/time.',
    'Primary actions: View Details; Row Again; Race Yourself when eligible; export; Done.':
        'Primary actions: Row Again; View Details; Done. Race Yourself remains available from Workout Details when the result is eligible.',
    'States and rules: Handle short, disconnected, and manually ended sessions without misleading summaries.':
        'States and rules: Distance leads with completion time; duration leads with distance; intervals use their program-specific work result and exclude Rest; Free Row leads with distance; Race shows the outcome against its fixed reference. Manually ended targets show actual progress and Ended Early. New Best compares only compatible completed workouts.',
}
for paragraph in doc.paragraphs:
    if paragraph.text == '13 Workout Complete':
        inside_complete = True
        continue
    if paragraph.text == '14 History':
        inside_complete = False
    if inside_complete and paragraph.text == 'PLANNED':
        paragraph.text = 'LOCKED'
        status_updated = True
    if inside_complete and paragraph.text in register_replacements:
        paragraph.text = register_replacements[paragraph.text]
if not status_updated:
    raise RuntimeError('Workout Complete status was not updated')

summary_replacements = {
    'Recommended next screens to define and lock: 1) Workout Complete, 2) History and Workout Details, 3) Connect Rower, 4) Settings and Display Customization.':
        'Recommended next screens to define and lock: 1) History and Workout Details, 2) Race Comparison Details, 3) Connect Rower, 4) Settings and Display Customization.',
    'Designed screens: Home, Quick Duration, Quick Distance, Quick Intervals, Segment Entry, Program Builder, My Programs, Workout Library, Program Detail / Start, Race Reference Picker, Live Row Free, Live Row Target, Live Row Race, and Pause / End are represented in the atlas.':
        'Designed screens: Home, Quick Duration, Quick Distance, Quick Intervals, Segment Entry, Program Builder, My Programs, Workout Library, Program Detail / Start, Race Reference Picker, Live Row Free, Live Row Target, Live Row Race, Pause / End, and Workout Complete are represented in the atlas.',
    'v2.0 - Consolidated all designed screens into an individual-state visual atlas, documented component changes by interaction, unified quick workout and program setup, simplified interval segments to Duration / Distance / Rest, incorporated Race Your Best across setup, reference selection, and Live Row, and locked the Pause / End safeguards.':
        'v2.0 - Consolidated all designed screens into an individual-state visual atlas, documented component changes by interaction, unified quick workout and program setup, simplified interval segments to Duration / Distance / Rest, incorporated Race Your Best across setup, reference selection, Live Row, and results, and locked Pause / End and Workout Complete.',
}
found = set()
for paragraph in doc.paragraphs:
    if paragraph.text in summary_replacements:
        original = paragraph.text
        paragraph.text = summary_replacements[original]
        found.add(original)
if found != set(summary_replacements):
    raise RuntimeError('Expected specification summary text was not found')

inventory = doc.inline_shapes[-1]
inventory_blip = inventory._inline.graphic.graphicData.pic.blipFill.blip
doc.part.related_parts[inventory_blip.embed]._blob = OVERVIEW.read_bytes()
inventory.width = Inches(7.65)
inventory.height = Inches(6.45)
for paragraph in doc.paragraphs:
    if paragraph.text.startswith('Figure 36. Complete Coxswain screen inventory:'):
        paragraph.text = 'Figure 42. Complete Coxswain screen inventory: 39 current screens and interaction states, plus 9 planned placeholders.'
        paragraph.paragraph_format.space_before = Pt(2)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.keep_together = True
        for run in paragraph.runs:
            run.font.size = Pt(7)
        break

anchor = next(p for p in doc.paragraphs if p.text == '17 v2.0 Consolidated Decisions')
inserted = []
for heading_text, filename, caption_text, bullets in STATES:
    page_break = doc.add_paragraph()
    page_break.add_run().add_break(WD_BREAK.PAGE)
    inserted.append(page_break)
    inserted.append(doc.add_paragraph(heading_text, style='Heading 2'))
    figure = doc.add_paragraph()
    figure.alignment = WD_ALIGN_PARAGRAPH.CENTER
    figure.add_run().add_picture(str(ROOT / 'work/screen-atlas' / filename), width=Inches(3.0))
    inserted.append(figure)
    caption = doc.add_paragraph(f'Current v2.0 state: {caption_text}')
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in caption.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(83, 100, 124)
    inserted.append(caption)
    inserted.append(doc.add_paragraph('Interaction changes', style='Heading 3'))
    for text in bullets:
        inserted.append(doc.add_paragraph(text, style='List Bullet'))

# Keep the consolidated decisions on a fresh page after the last screen state.
final_break = doc.add_paragraph()
final_break.add_run().add_break(WD_BREAK.PAGE)
inserted.append(final_break)
for paragraph in inserted:
    anchor._p.addprevious(paragraph._p)

doc.save(SPEC)
print('added Workout Complete states to the specification')
