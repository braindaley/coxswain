from pathlib import Path
from docx import Document
from docx.shared import RGBColor

SPEC = Path('/Users/brian/Documents/ChatGPT/Coxswain/outputs/Coxswain_Product_Design_Spec.docx')
doc = Document(SPEC)

def replace_paragraph(paragraph, old, new):
    if old not in paragraph.text:
        return
    text = paragraph.text.replace(old, new)
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ''
    else:
        paragraph.add_run(text)

for paragraph in doc.paragraphs:
    replace_paragraph(paragraph, 'v1.4', 'v1.5')

for section in doc.sections:
    for paragraph in section.header.paragraphs:
        replace_paragraph(paragraph, 'v1.4', 'v1.5')

for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                replace_paragraph(paragraph, 'v1.4', 'v1.5')

current_screen = None
for paragraph in doc.paragraphs:
    if paragraph.text.startswith('06 Program Builder'):
        current_screen = 'builder'
    elif paragraph.text.startswith('07 Workout Library'):
        current_screen = 'library'
    elif paragraph.style.name.startswith('Heading 2') and current_screen:
        current_screen = None
    elif current_screen in {'builder', 'library'} and paragraph.text == 'PLANNED':
        paragraph.text = 'IN PROGRESS'
        if paragraph.runs:
            paragraph.runs[0].bold = True
            paragraph.runs[0].font.color.rgb = RGBColor(11, 99, 246)
    elif current_screen == 'library' and paragraph.text.startswith('Primary actions:'):
        paragraph.text = 'Primary actions: View Program; Start; optionally save a copy to My Programs. The library has no Create Program action.'
    elif current_screen == 'library' and paragraph.text.startswith('States and rules:'):
        paragraph.text = 'States and rules: Library content is read-only. Use the same compact card system as My Programs, omit the Create Program button, and omit owned-program management actions. Empty and loading states depend on how content is packaged.'

doc.add_page_break()
doc.add_heading('12 Workout Library and Program Builder Direction', level=1)
doc.add_heading('Workout Library', level=2)
doc.add_paragraph('Workout Library uses the same compact list and card components as My Programs. The tab, card density, type-specific one-line summaries, View Program action, Start action, and Programs navigation remain consistent.')
doc.add_paragraph('Because the library is a curated read-only collection, it does not show Create Program. Library cards also omit Edit, Duplicate, Export, and Delete. A future Save to My Programs action may create an owned copy without changing the library item.')
doc.add_heading('Program Builder', level=2)
doc.add_paragraph('Program Builder reuses the approved Duration and Distance setup components: large direct target entry, familiar preset choices, optional goal selection, goal target entry, and an always-visible summary. It adds a program name and replaces Start Workout with Save Program.')
doc.add_paragraph('The type selector supports Duration, Distance, and Intervals. Duration and Distance display their existing target and preset controls. Intervals replace the single target with an ordered segment list. Each segment has a Work or Rest role and a duration or distance value; users can add, edit, reorder, or remove segments.')
doc.add_paragraph('Save Program validates a non-empty name and a complete configuration, then returns the saved definition to My Programs. Back with unsaved changes asks whether to discard them. Saving a program does not start a workout.')
doc.add_heading('Version update', level=2)
doc.add_paragraph('v1.5 - Applied the compact Programs card system to Workout Library without a Create Program action. Added the Program Builder design direction using the approved Duration and Distance setup components, with interval segment support.')

doc.save(SPEC)
