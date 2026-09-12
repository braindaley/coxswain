from pathlib import Path
from docx import Document
from docx.enum.text import WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

SPEC = Path('/Users/brian/Documents/ChatGPT/Coxswain/outputs/Coxswain_Product_Design_Spec.docx')

def replace_text(paragraph, replacements):
    original = paragraph.text
    updated = original
    for old, new in replacements:
        updated = updated.replace(old, new)
    if updated != original:
        if paragraph.runs:
            paragraph.runs[0].text = updated
            for run in paragraph.runs[1:]:
                run.text = ''
        else:
            paragraph.add_run(updated)

def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tc_pr.append(shd)
    shd.set(qn('w:fill'), fill)

def border_table(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn('w:tblBorders'))
    if borders is None:
        borders = OxmlElement('w:tblBorders')
        tbl_pr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        tag = qn(f'w:{edge}')
        node = borders.find(tag)
        if node is None:
            node = OxmlElement(f'w:{edge}')
            borders.append(node)
        node.set(qn('w:val'), 'single')
        node.set(qn('w:sz'), '4')
        node.set(qn('w:color'), 'D9D9D9')

doc = Document(SPEC)

paragraph_replacements = [
    ('| v1.3', '| v1.4'),
    ('Quick Start contains four actions: Duration, Distance, Intervals, Workouts.', 'Quick Start contains four actions: Duration, Distance, My Programs, Library.'),
    ('Bottom navigation remains Home, Workouts, History, More.', 'Bottom navigation remains Home, Programs, History, More.'),
    ('Bottom navigation: Home, Workouts, History, More.', 'Bottom navigation: Home, Programs, History, More.'),
    ('Home, Workouts, History, and More are the persistent bottom-level destinations.', 'Home, Programs, History, and More are the persistent bottom-level destinations.'),
    ('Duration, Distance, Intervals, Workouts quick actions.', 'Duration, Distance, My Programs, Library quick actions.'),
    ('Duration, Distance, Intervals, Workouts. 4 evenly distributed actions.', 'Duration, Distance, My Programs, Library. Four evenly distributed actions.'),
    ('Home, Workouts, History, More. Active item uses primary color.', 'Home, Programs, History, More. Active item uses primary color.'),
    ('Four fixed destinations.', 'Four fixed destinations: Home, Programs, History, and More.'),
    ('Locked items in v1.2', 'Locked items in v1.4'),
    ('Recommended next screens to define and lock: 1) Intervals List and Builder, 2) Premade Workouts Library, 3) Live Row (standard), 4) Live Row goal states, 5) History and Workout Details.', 'Recommended next screens to define and lock: 1) Program Builder, 2) Workout Library, 3) Program Detail / Start, 4) Live Row Target and goal states, 5) History and Workout Details.'),
]

for p in doc.paragraphs:
    replace_text(p, paragraph_replacements)

for section in doc.sections:
    for p in section.header.paragraphs:
        replace_text(p, [('v1.1', 'v1.4')])

for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                replace_text(p, paragraph_replacements)

        cells = row.cells
        if len(cells) >= 4:
            control = cells[0].text.strip()
            if control == 'Intervals':
                cells[0].text = 'My Programs'
                cells[1].text = 'Open saved reusable rowing programs.'
                cells[2].text = 'programs'
                cells[3].text = 'Duration, distance, and interval templates that can be viewed, edited, or started.'
            elif control == 'Workouts':
                cells[0].text = 'Library'
                cells[1].text = 'Open the curated Workout Library.'
                cells[2].text = 'programs/library'
                cells[3].text = 'Premade reusable programs supplied by the product owner.'
            elif control == 'Workouts nav':
                cells[0].text = 'Programs nav'
                cells[1].text = 'Open Programs with My Programs and Workout Library.'
                cells[2].text = 'programs'
                cells[3].text = 'Reusable plans are kept separate from completed workout history.'

        if len(cells) >= 4 and cells[0].text.strip().isdigit():
            number = cells[0].text.strip()
            mapping = {
                '05': ('My Programs', 'IN PROGRESS', 'Saved reusable duration, distance, and interval templates.'),
                '06': ('Program Builder', 'Planned', 'Create or edit a reusable duration, distance, or interval program.'),
                '07': ('Workout Library', 'Planned', 'Curated reusable programs supplied by the product owner.'),
                '08': ('Program Detail / Start', 'Planned', 'Review a reusable program and start a new workout session.'),
            }
            if number in mapping:
                cells[1].text, cells[2].text, cells[3].text = mapping[number]

doc.add_page_break()
doc.add_heading('09 Terminology and Information Architecture', level=1)
doc.add_paragraph('This section defines the product language and the complete screen contract for v1.4. It supersedes older navigation labels shown in retained v1.0-v1.2 mockup images.')
doc.add_heading('Program and workout definitions', level=2)
doc.add_paragraph('Program: a reusable rowing template. A program may contain a duration target, a distance target, or an ordered set of work and rest intervals. It may also contain an optional performance goal. Programs can be started repeatedly without changing the saved definition.')
doc.add_paragraph('Workout: an active or completed rowing session. Starting a program creates a workout. History contains completed workouts and their measured results.')
doc.add_heading('Primary navigation', level=2)
nav = doc.add_table(rows=1, cols=3)
nav.alignment = WD_TABLE_ALIGNMENT.CENTER
nav.autofit = False
for i, text in enumerate(('Destination', 'Contains', 'Purpose')):
    nav.rows[0].cells[i].text = text
for values in [
    ('Home', 'Connection, last workout, Free Row, Quick Start', 'Start common rowing flows quickly.'),
    ('Programs', 'My Programs and Workout Library', 'Find, create, manage, and start reusable plans.'),
    ('History', 'Completed workouts and results', 'Review, compare, export, repeat, or race past sessions.'),
    ('More', 'Settings, connection, data, diagnostics, about', 'Access secondary tools and configuration.'),
]:
    row = nav.add_row().cells
    for i, value in enumerate(values):
        row[i].text = value
for row_index, row in enumerate(nav.rows):
    for cell in row.cells:
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shade(cell, '10213F' if row_index == 0 else ('F4F7FB' if row_index % 2 == 0 else 'FFFFFF'))
        for p in cell.paragraphs:
            for run in p.runs:
                run.font.color.rgb = RGBColor(255,255,255) if row_index == 0 else RGBColor(16,33,63)
                run.font.size = Pt(9)
                run.font.bold = row_index == 0
border_table(nav)

doc.add_heading('Home launcher revision', level=2)
doc.add_paragraph('Home Quick Start uses Duration, Distance, My Programs, and Library. My Programs opens the saved-template list within Programs. Library opens the curated Workout Library tab. The persistent navigation uses Home, Programs, History, and More.')

doc.add_page_break()
doc.add_heading('10 Complete Screen Specification Register', level=1)
doc.add_paragraph('Every screen below has a defined role, primary content, actions, and state requirements. LOCKED means the design has been approved. IN PROGRESS means the current direction is being reviewed. PLANNED screens have a functional contract but still need visual approval.')

screens = [
('01 Home','LOCKED','Launch rowing modes and show connection and recent activity.','Coxswain brand, connection card, hero image, Last Workout with Row Again, Free Row, Quick Start.','Settings; connection details; workout details; Row Again; Free Row; Duration; Distance; My Programs; Library.','Disconnected start actions show the app-level connection prompt. Persistent navigation: Home, Programs, History, More.'),
('02 Live Row Free','LOCKED','Run an open-ended rowing session.','Equal-area configurable metric slots on #042C3D. All values use one maximized shared size.','Edit display; replace/add/remove metrics; Pause/Resume; End session with confirmation.','No bottom navigation. Preserve at least one metric. Connection-loss and missing-sensor states remain to be finalized.'),
('03 Duration Setup','LOCKED','Configure a timed workout.','Direct duration entry, presets, optional goal, live Workout Summary.','Choose duration; choose None, Stroke Rate, Speed, or Power; set target; Start Workout.','Connection, audio, and display controls remain outside setup. Starting creates a workout session.'),
('04 Distance Setup','LOCKED','Configure a target-distance workout.','Direct distance entry, presets, optional goal, live Workout Summary.','Choose distance; choose None, Stroke Rate, Speed, or Power; set target; Start Workout.','Uses the Duration Setup interaction pattern. Pace-boat selection remains outside this screen.'),
('05 My Programs','IN PROGRESS','Manage saved reusable programs.','Compact cards for duration, distance, and interval programs, with type-specific summaries.','Create Program; View Program; Start; Edit; Duplicate; Export; Delete.','Delete requires confirmation. Empty state invites the first program. Programs navigation is selected.'),
('06 Program Builder','PLANNED','Create or edit a reusable program.','Program name, type choice, configuration, optional goals, interval segment list when applicable.','Save; add/edit/reorder/remove segments; preview summary; cancel.','Supports duration, distance, and interval definitions. Validate mixed-unit and empty-segment cases.'),
('07 Workout Library','PLANNED','Browse curated reusable programs.','Categories or grouped list, concise program cards, description and difficulty metadata supplied by the product owner.','View; start; optionally save a copy to My Programs.','Library content is read-only. Empty/loading states depend on how content is packaged.'),
('08 Program Detail Start','PLANNED','Review a program before starting a workout.','Name, description, full configuration, ordered segments, goal, estimated totals when meaningful.','Start Workout; Edit for owned programs; save copy for library programs; back.','Do not present an exact total duration when distance-based segments make it unknowable.'),
('09 Live Row Target','PLANNED','Run duration, distance, interval, or library workouts.','The approved configurable metric grid plus target/progress context.','Edit display; Pause/Resume; End session.','Show current performance relative to the selected goal without reducing metric readability.'),
('10 Race Yourself Setup','PLANNED','Choose a completed workout as the pace boat.','Compatible prior workouts, date, result, and comparison eligibility.','Select workout; review; start race; cancel.','Filter to comparable workout structures and explain why an item is unavailable.'),
('11 Live Row Race Yourself','PLANNED','Race against snapshot data from a prior workout.','Configurable metrics with prominent distance and time deltas.','Edit display; Pause/Resume; End session.','Use separate saved metric configuration. Ahead/behind meaning must use text or sign as well as color.'),
('12 Pause End','PLANNED','Safely pause, resume, finish, or discard an active workout.','Paused state, elapsed status, resume action, end choices.','Resume; End and Save; Discard Workout; cancel.','Discard requires explicit confirmation. End and Save proceeds to Workout Complete.'),
('13 Workout Complete','PLANNED','Confirm completion and summarize results.','Key measured totals, achieved target, interval results, date/time.','View Details; Row Again; Race Yourself when eligible; export; Done.','Handle short, disconnected, and manually ended sessions without misleading summaries.'),
('14 History','PLANNED','Browse completed workouts.','Chronological workout list, filters, totals, and search if needed.','Open details; repeat; race; filter.','History contains executed workouts, never reusable program definitions.'),
('15 Workout Details','PLANNED','Inspect one completed workout.','Metrics, graphs, interval breakdown, equipment/source metadata.','Export; Row Again; Race Yourself; delete with confirmation.','Use real recorded values and make unavailable sensor series explicit.'),
('16 Race Comparison Details','PLANNED','Compare a completed workout with its pace workout.','Result summary, distance/time deltas, aligned splits or snapshots.','Switch compared metrics; open either workout; repeat race.','Keep comparison scales and signs consistent. Explain incomplete comparisons.'),
('17 Connect Rower','PLANNED','Connect and inspect WaterRower hardware.','USB and Bluetooth status, device name, transport, troubleshooting state.','Scan/connect; disconnect; retry; open diagnostics.','UI consumes repository/session state rather than handling protocols directly.'),
('18 Settings','PLANNED','Set global app behavior.','Units, live-display defaults, audio feedback, heart-rate source, data, advanced options.','Change settings; open sub-screens; reset applicable defaults.','Settings apply consistently and should not clutter workout setup.'),
('19 Display Customization','PLANNED','Choose and order live-row metrics.','Available metrics and ordered active slots for normal and race modes.','Add; remove; replace; reorder; preview; restore defaults.','At least one metric is required. The shared live-number size recalculates from the final set.'),
('20 Heart Rate Source','PLANNED','Choose the source of pulse data.','WaterRower, Android sensor, Bluetooth LE, and ANT+ options when supported.','Select; connect/authorize as needed; test reading.','Show unavailable hardware and permission states clearly.'),
('21 Data Export','PLANNED','Export, back up, and manage stored data.','TCX, calendar, supported integrations, backup/storage information.','Export selected workout; export archive; import/restore if supported; storage management.','Confirm destructive storage actions and show success/failure details.'),
('22 More Diagnostics','PLANNED','Collect secondary tools in one place.','Connection, settings, data/export, diagnostics, protocol info, about.','Open each destination; copy diagnostic information where appropriate.','Keep this screen concise; detailed controls live on destination screens.'),
]

for title, status, purpose, content, actions, states in screens:
    doc.add_heading(title, level=2)
    p = doc.add_paragraph()
    r = p.add_run(status)
    r.bold = True
    r.font.color.rgb = RGBColor(11,99,246) if status != 'PLANNED' else RGBColor(83,100,124)
    doc.add_paragraph(f'Purpose: {purpose}')
    doc.add_paragraph(f'Primary content: {content}')
    doc.add_paragraph(f'Primary actions: {actions}')
    doc.add_paragraph(f'States and rules: {states}')

doc.add_page_break()
doc.add_heading('11 My Programs Design Direction', level=1)
doc.add_paragraph('The Programs destination contains two tabs: My Programs and Workout Library. My Programs is selected in the current design. It uses compact cards because users may save many programs.')
doc.add_heading('My Programs card content', level=2)
doc.add_paragraph('Each card shows a user-defined name and a one-line summary suited to its type. Examples include “60 min · No goal,” “5,000 m · Target 26 SPM,” and “6 × 500 m · 1 min rest.” The card does not assume every program is an interval workout.')
doc.add_heading('Actions and behavior', level=2)
doc.add_paragraph('Create Program begins a reusable duration, distance, or interval definition. View Program opens Program Detail / Start. Start creates a new workout from the saved program. The overflow menu contains Edit, Duplicate, Export, and Delete. Delete requires confirmation.')
doc.add_heading('Layout and navigation', level=2)
doc.add_paragraph('Use a plain #F4F7FB header without a scenic image. Compact white cards use 20dp corners, restrained elevation, 48dp action targets, and small vertical gaps. The bottom navigation shows Home, Programs, History, and More, with Programs selected.')
doc.add_heading('Version update', level=2)
doc.add_paragraph('v1.4 - Replaced ambiguous reusable-workout terminology with Programs. Added My Programs and Workout Library under the Programs destination, revised Home Quick Start and primary navigation labels, and added functional specifications for all 22 screens.')

doc.save(SPEC)
