from pathlib import Path
from docx import Document

SPEC = Path('/Users/brian/Documents/ChatGPT/Coxswain/outputs/Coxswain_Product_Design_Spec.docx')
doc = Document(SPEC)

def replace_text(paragraph, old, new):
    if old not in paragraph.text:
        return
    updated = paragraph.text.replace(old, new)
    if paragraph.runs:
        paragraph.runs[0].text = updated
        for run in paragraph.runs[1:]:
            run.text = ''
    else:
        paragraph.add_run(updated)

replacements = [
    ('v1.6', 'v1.7'),
    ('500 m, 1,000 m, 5,000 m, 10,000 m, Half Marathon, and Full Marathon', '500 m, 1,000 m, 2,000 m, 5,000 m, 6,000 m, and 10,000 m'),
    ('500 m, 1,000 m, 5,000 m, 10,000 m, Half Marathon, Full Marathon', '500 m, 1,000 m, 2,000 m, 5,000 m, 6,000 m, 10,000 m'),
    ('500 m, 1,000 m, 5,000 m, 10,000 m, Half Marathon, and Full Marathon.', '500 m, 1,000 m, 2,000 m, 5,000 m, 6,000 m, and 10,000 m.'),
]
for paragraph in doc.paragraphs:
    for old, new in replacements:
        replace_text(paragraph, old, new)
for section in doc.sections:
    for paragraph in section.header.paragraphs:
        for old, new in replacements:
            replace_text(paragraph, old, new)
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for old, new in replacements:
                    replace_text(paragraph, old, new)

if not any(p.text == '14 Shared Set Goal and Distance Presets' for p in doc.paragraphs):
    doc.add_page_break()
    doc.add_heading('14 Shared Set Goal and Distance Presets', level=1)
    doc.add_heading('Shared Set Goal component', level=2)
    doc.add_paragraph('Program Builder uses the approved Distance Workout Set Goal component instead of a separate simplified control. Duration Setup, Distance Setup, and Program Builder therefore share the same four icon tiles: None, Stroke Rate, Speed, and Power.')
    doc.add_paragraph('Selecting a performance goal reveals its target panel. The panel provides minus and plus controls, an appropriate slider, and a tappable central value for direct entry. The selected tile uses the Material 3 selected state, and every change immediately updates the summary.')
    doc.add_heading('Rowing distance presets', level=2)
    doc.add_paragraph('The approved distance presets are 500 m, 1,000 m, 2,000 m, 5,000 m, 6,000 m, and 10,000 m. Half Marathon and Full Marathon are removed because their running terminology does not fit the rowing product. Direct entry continues to support any distance outside these presets.')
    doc.add_paragraph('This written v1.7 rule supersedes the Half Marathon and Full Marathon labels visible in the retained earlier Distance Setup mockup image.')
    doc.add_heading('Version update', level=2)
    doc.add_paragraph('v1.7 - Standardized Program Builder on the approved Set Goal component and replaced running-oriented marathon presets with rowing distances.')

doc.save(SPEC)
