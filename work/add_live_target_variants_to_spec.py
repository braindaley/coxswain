from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches


SPEC = Path('/Users/brian/Documents/ChatGPT/Coxswain/outputs/Coxswain_Product_Design_Spec.docx')
ATLAS = Path('/Users/brian/Documents/ChatGPT/Coxswain/work/screen-atlas')
doc = Document(SPEC)

replacements = {
    'Primary content: The approved configurable metric grid plus target/progress context.':
        'Primary content: The approved configurable metric grid plus a compact progress panel. Duration shows elapsed and remaining time; distance shows completed and remaining meters; intervals show the current ordered segment, its progress, and the next segment.',
    'States and rules: Show current performance relative to the selected goal without reducing metric readability.':
        'States and rules: The metric grid, Edit display interaction, Pause, and End session match Live Row Free. The progress panel is fixed below the grid. Free-form intervals advance through their saved order and never introduce rounds.',
    'Designed screens: Home, Quick Duration, Quick Distance, Quick Intervals, Segment Entry, Program Builder, My Programs, Workout Library, Program Detail / Start, Live Row Free, and Live Row Race are represented in the atlas.':
        'Designed screens: Home, Quick Duration, Quick Distance, Quick Intervals, Segment Entry, Program Builder, My Programs, Workout Library, Program Detail / Start, Live Row Free, Live Row Target, and Live Row Race are represented in the atlas.',
    'Figure 26. Complete Coxswain screen inventory: 22 current screens and interaction states, plus 13 planned placeholders.':
        'Figure 29. Complete Coxswain screen inventory: 25 current screens and interaction states, plus 12 planned placeholders.',
}

found = set()
inside_target = False
for paragraph in doc.paragraphs:
    text = paragraph.text.strip()
    if text == '09 Live Row Target':
        inside_target = True
        continue
    if inside_target and text == 'PLANNED':
        paragraph.text = 'LOCKED'
        inside_target = False
        continue
    if text in replacements:
        paragraph.text = replacements[text]
        found.add(text)
missing = set(replacements) - found
if missing:
    raise RuntimeError('Expected text was not found: ' + repr(sorted(missing)))

anchor = next(p for p in doc.paragraphs if p.text.strip() == '17 v2.0 Consolidated Decisions')

def add_variant(number, title, filename, bullets):
    heading_text = f'16.{number} Live Row Target — {title}'
    if any(p.text.strip() == heading_text for p in doc.paragraphs):
        return
    heading = anchor.insert_paragraph_before(heading_text, style='Heading 2')
    heading.paragraph_format.page_break_before = True
    picture = anchor.insert_paragraph_before('', style='Normal')
    picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture.add_run().add_picture(str(ATLAS / filename), width=Inches(3))
    caption = anchor.insert_paragraph_before('', style='Normal')
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = caption.add_run(f'Current v2.0 state: Live Row Target — {title}')
    run.italic = True
    anchor.insert_paragraph_before('Interaction changes', style='Heading 3')
    for bullet in bullets:
        anchor.insert_paragraph_before(bullet, style='List Bullet')
    anchor.insert_paragraph_before('', style='Normal')

add_variant(23, 'Duration', '23_live_target_duration.png', [
    'The metric grid remains unchanged while progress shows elapsed time, remaining time, and the fixed target.',
    'The progress panel does not consume a configurable metric position.',
])
add_variant(24, 'Distance', '24_live_target_distance.png', [
    'Progress shows completed meters, remaining meters, and the fixed distance target.',
    'All configurable metric values retain one shared maximum font size.',
])
add_variant(25, 'Intervals', '25_live_target_intervals.png', [
    'Progress names the current Row or Rest segment, its remaining target, and the next ordered segment.',
    'The screen follows the builder’s free-form segment order and has no rounds model.',
])

doc.save(SPEC)
print('added duration, distance, and interval Live Row Target variants')
