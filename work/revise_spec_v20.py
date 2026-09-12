from pathlib import Path
import json
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

SPEC = Path('/Users/brian/Documents/ChatGPT/Coxswain/outputs/Coxswain_Product_Design_Spec.docx')
ATLAS = Path('/Users/brian/Documents/ChatGPT/Coxswain/work/screen-atlas')
doc = Document(SPEC)
if len(doc.sections) == 0:
    doc._element.body.get_or_add_sectPr()

def replace_text(paragraph, replacements):
    original = paragraph.text
    updated = original
    for old, new in replacements:
        updated = updated.replace(old, new)
    if updated == original:
        return
    if paragraph.runs:
        paragraph.runs[0].text = updated
        for run in paragraph.runs[1:]:
            run.text = ''
    else:
        paragraph.add_run(updated)

replacements = [
    ('v1.8', 'v2.0'),
    ('10 Race Yourself Setup', '10 Race Reference Picker'),
    ('Choose a completed workout as the pace boat.', 'Choose an alternative compatible completed workout for Race Your Best.'),
    ('The type selector supports Duration, Distance, and Intervals. Duration and Distance display their existing target and preset controls. Intervals replace the single target with an ordered segment list. Each segment has a Work or Rest role and a duration or distance value; users can add, edit, reorder, or remove segments.',
     'The type selector supports Duration, Distance, and Intervals. Duration and Distance display their existing target and preset controls. Intervals replace the single target with an ordered segment list. Each segment is exactly Duration, Distance, or Rest and reuses the corresponding large target-entry and preset component. Users can add, edit, reorder, or delete segments.'),
]
for p in doc.paragraphs:
    replace_text(p, replacements)
for section in doc.sections:
    for p in section.header.paragraphs:
        replace_text(p, replacements)
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                replace_text(p, replacements)

# Update inventory row for the race-reference flow.
for table in doc.tables:
    for row in table.rows:
        cells = row.cells
        if len(cells) >= 4 and cells[0].text.strip() == '10':
            cells[1].text = 'Race Reference Picker'
            cells[2].text = 'Planned'
            cells[3].text = 'Choose another compatible completed workout; Workout Setup defaults to the best result.'

items = json.loads((ATLAS / 'manifest.json').read_text(encoding='utf-8'))

doc.add_page_break()
doc.add_heading('16 Current Visual Screen Atlas', level=1)
doc.add_paragraph('This section is the current visual source of truth for every screen and interaction state designed through v2.0. Each state has its own screenshot so reviewers can assess component changes without interpreting a composite image. Earlier composite mockups remain in the document as design history; where labels or controls differ, this atlas supersedes them.')
doc.add_heading('Shared interaction model', level=2)
for text in [
    'Quick Duration, Distance, and Intervals use one Workout Setup shell. Only the target configuration area changes by type.',
    'Quick setup keeps Start Workout primary and Save as Program secondary. Program creation adds Program Name and changes the primary action to Save Program.',
    'Set Goal uses the same None, Stroke Rate, Speed, and Power tiles. Selecting a goal reveals its value controls and updates Summary immediately.',
    'Interval segments use three choices: Duration, Distance, and Rest. Each choice reuses the large value-entry card and relevant presets. Existing rows open Edit with Save and Delete.',
    'Race Your Best is selected in Distance Workout Setup and replaces Set Goal while active. Live Row displays the fixed ahead/behind comparison.',
]:
    doc.add_paragraph(text, style='List Bullet')

matrix = doc.add_table(rows=1, cols=4)
matrix.alignment = WD_TABLE_ALIGNMENT.CENTER
matrix.autofit = False
headers = ('Entry context','Program name','Primary action','Secondary action')
for i,h in enumerate(headers): matrix.rows[0].cells[i].text=h
rows = [
    ('Quick workout','Hidden','Start Workout','Save as Program'),
    ('Create program','Visible','Save Program','Cancel / Back'),
    ('Edit program','Visible','Save Changes','Cancel / Back'),
    ('Library program','Read-only','Start Workout','Save to My Programs'),
]
for values in rows:
    cells=matrix.add_row().cells
    for i,value in enumerate(values): cells[i].text=value
for ri,row in enumerate(matrix.rows):
    for cell in row.cells:
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        props=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),'10213F' if ri==0 else ('F4F7FB' if ri%2==0 else 'FFFFFF')); props.append(shd)
        for p in cell.paragraphs:
            for run in p.runs:
                run.font.size=Pt(8.5); run.font.bold=ri==0; run.font.color.rgb=RGBColor(255,255,255) if ri==0 else RGBColor(16,33,63)

for number,(title,filename,notes) in enumerate(items,1):
    doc.add_page_break()
    doc.add_heading(f'16.{number:02d} {title}', level=2)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(ATLAS/filename),width=Inches(3.0))
    caption=doc.add_paragraph(f'Current v2.0 state: {title}')
    caption.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for run in caption.runs:
        run.italic=True; run.font.size=Pt(8); run.font.color.rgb=RGBColor(83,100,124)
    doc.add_heading('Interaction changes', level=3)
    for note in notes:
        p=doc.add_paragraph(note,style='List Bullet')
        p.paragraph_format.space_after=Pt(2)

doc.add_page_break()
doc.add_heading('17 v2.0 Consolidated Decisions', level=1)
decisions = [
    ('Designed screens', 'Home, Quick Duration, Quick Distance, Quick Intervals, Segment Entry, Program Builder, My Programs, Workout Library, Live Row Free, and Live Row Race are represented in the atlas.'),
    ('Rowing distances', 'Presets are 500 m, 1K, 2K, 5K, 6K, and 10K. Direct entry supports other distances.'),
    ('Segment model', 'A segment is Duration, Distance, or Rest. Add and Edit share the same component; Edit adds Delete.'),
    ('Race placement', 'Selection belongs in Workout Setup. Live Row displays comparison. Race Your Best and Set Goal are mutually exclusive.'),
    ('Terminology', 'Programs are reusable definitions. Workouts are active or completed sessions. Quick setup therefore uses Save as Program.'),
    ('Supersession', 'The v2.0 atlas replaces obsolete Quick Start labels, marathon presets, separate goal controls, and earlier interval-segment forms shown in retained figures.'),
]
for label,value in decisions:
    p=doc.add_paragraph(style='List Bullet'); r=p.add_run(label+': '); r.bold=True; p.add_run(value)
doc.add_heading('Version update', level=2)
doc.add_paragraph('v2.0 - Consolidated all designed screens into an individual-state visual atlas, documented component changes by interaction, unified quick workout and program setup, simplified interval segments to Duration / Distance / Rest, and incorporated Race Your Best across setup and Live Row.')

doc.save(SPEC)
print(f'updated {SPEC} with {len(items)} atlas screenshots')
