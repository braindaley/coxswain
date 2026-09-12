from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches


SPEC = Path('/Users/brian/Documents/ChatGPT/Coxswain/outputs/Coxswain_Product_Design_Spec.docx')
ATLAS = Path('/Users/brian/Documents/ChatGPT/Coxswain/work/screen-atlas')
doc = Document(SPEC)

replacements = {
    'Primary content: The approved configurable metric grid plus a compact progress panel. Duration shows elapsed and remaining time; distance shows completed and remaining meters; intervals show the complete saved Row / Rest sequence on one line with the active segment highlighted.':
        'Primary content: The approved configurable metric grid plus a compact progress panel. Duration shows elapsed and remaining time; distance shows completed and remaining meters; intervals show the complete saved sequence as one continuous segmented line.',
    'States and rules: The metric grid, Edit display interaction, Pause, and End session match Live Row Free. The progress panel is fixed below the grid. Free-form intervals advance through their saved order and never introduce rounds. When Stroke rate, Speed, or Power has a goal, its metric card keeps the current reading as the large value and shows a compact signed variance in the upper-right. Negative variance is red, positive variance is green, and exact target is neutral.':
        'States and rules: The metric grid, Pause, and End session match Live Row Free during Row segments. The progress panel is fixed below the grid. Free-form intervals advance through their saved order and never introduce rounds. During Rest, the metric grid becomes a large countdown and identifies the next Row segment. When Stroke rate, Speed, or Power has a goal, its metric card keeps the current reading as the large value and shows a small signed variance in the upper-right. Negative variance is red, positive variance is green, and exact target is neutral.',
    'The complete saved sequence appears on one line as Row, Rest, Row, Rest, and the active segment is highlighted.':
        'One continuous segmented line represents the complete saved sequence. Blue sections are Row and lighter sections are Rest.',
    'The active segment shows its remaining target. The screen follows the builder’s free-form order and has no rounds model.':
        'A small marker shows current progress. The heading identifies the active segment and remaining time; the line does not repeat segment labels or durations.',
    'Live performance goals: Stroke rate, Speed, and Power keep their current reading as the large card value. A compact signed variance appears in the upper-right of that same card: red below target, green above target, and neutral on target. Variance is never a separate metric card.':
        'Live performance goals: Stroke rate, Speed, and Power keep their current reading as the large card value. A small signed number appears in the upper-right of that same card: red below target, green above target, and neutral on target. Variance is never a separate metric card.',
    'A small red −2 SPM badge in the upper-right shows that the current reading is below target.':
        'A small red −2 SPM value in the upper-right shows that the current reading is below target.',
    'A small signed variance badge uses the same upper-right position; −0.7 is red because current speed is below target.':
        'A small signed variance uses the same upper-right position; −0.7 is red because current speed is below target.',
    'A green +8 W badge in the upper-right shows that current power is above target.':
        'A small green +8 W value in the upper-right shows that current power is above target.',
    'Figure 32. Complete Coxswain screen inventory: 28 current screens and interaction states, plus 12 planned placeholders.':
        'Figure 33. Complete Coxswain screen inventory: 29 current screens and interaction states, plus 12 planned placeholders.',
}

found = set()
for paragraph in doc.paragraphs:
    text = paragraph.text.strip()
    if text in replacements:
        paragraph.text = replacements[text]
        found.add(text)
missing = set(replacements) - found
if missing:
    raise RuntimeError('Expected text was not found: ' + repr(sorted(missing)))

anchor = next(p for p in doc.paragraphs if p.text.strip() == '17 v2.0 Consolidated Decisions')
heading_text = '16.29 Live Row Target — Interval Rest'
if not any(p.text.strip() == heading_text for p in doc.paragraphs):
    heading = anchor.insert_paragraph_before(heading_text, style='Heading 2')
    heading.paragraph_format.page_break_before = True
    picture = anchor.insert_paragraph_before('', style='Normal')
    picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture.add_run().add_picture(str(ATLAS / '29_live_target_interval_rest.png'), width=Inches(3))
    caption = anchor.insert_paragraph_before('', style='Normal')
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = caption.add_run('Current v2.0 state: Live Row Target — Interval Rest')
    run.italic = True
    anchor.insert_paragraph_before('Interaction changes', style='Heading 3')
    anchor.insert_paragraph_before('When a Rest segment begins, the configurable metric grid becomes a large countdown for the remaining rest time.', style='List Bullet')
    anchor.insert_paragraph_before('The next Row segment, segmented sequence line, current-position marker, Pause, and End session remain visible.', style='List Bullet')
    anchor.insert_paragraph_before('', style='Normal')

doc.save(SPEC)
print('updated interval line and added interval Rest state')
