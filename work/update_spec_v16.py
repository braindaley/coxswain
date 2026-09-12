from pathlib import Path
from docx import Document
from docx.shared import RGBColor

SPEC = Path('/Users/brian/Documents/ChatGPT/Coxswain/outputs/Coxswain_Product_Design_Spec.docx')
doc = Document(SPEC)

def replace_text(paragraph, old, new):
    if old not in paragraph.text:
        return
    updated = paragraph.text.replace(old, new)
    if paragraph.runs:
        paragraph.runs[0].text = updated
        for run in paragraph.runs[1:]:
            run.text = ''
    else:
        paragraph.add_run(updated)

for paragraph in doc.paragraphs:
    replace_text(paragraph, 'v1.5', 'v1.6')
for section in doc.sections:
    for paragraph in section.header.paragraphs:
        replace_text(paragraph, 'v1.5', 'v1.6')
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                replace_text(paragraph, 'v1.5', 'v1.6')

current_screen = None
for paragraph in doc.paragraphs:
    if paragraph.text.startswith('05 My Programs'):
        current_screen = 'my_programs'
    elif paragraph.style.name.startswith('Heading 2') and current_screen:
        current_screen = None
    elif current_screen == 'my_programs' and paragraph.text == 'IN PROGRESS':
        paragraph.text = 'LOCKED'
        paragraph.runs[0].bold = True
        paragraph.runs[0].font.color.rgb = RGBColor(11, 99, 246)
    elif current_screen == 'my_programs' and paragraph.text.startswith('States and rules:'):
        paragraph.text = ('States and rules: Delete requires confirmation. Empty state explains programs and offers Create Program. '
                          'Order programs by most recently used. Tapping the card name opens Program Detail; Start remains explicit. '
                          'Show search and filters when the collection reaches 10 programs. Programs navigation is selected.')

if not any(p.text == 'Approved My Programs decisions' for p in doc.paragraphs):
    doc.add_page_break()
    doc.add_heading('13 Approved My Programs Decisions', level=1)
    doc.add_paragraph('The My Programs screen and the decisions in this section are approved and locked.')
    decisions = [
        ('Ordering', 'Sort by most recently used, with newly created programs treated as most recent.'),
        ('Card interaction', 'Tapping the program name or View Program opens Program Detail. Start remains a separate explicit action.'),
        ('Duplicate', 'Duplicate appends “copy” to the name and immediately opens the new owned program for editing.'),
        ('Library ownership', 'Workout Library items remain read-only. Save to My Programs creates a separate editable copy.'),
        ('Large collections', 'Keep the default screen compact. Show search and type filters when My Programs contains 10 or more items.'),
        ('Empty state', 'Explain that programs are reusable duration, distance, or interval templates and provide a prominent Create Program action.'),
    ]
    for label, rule in decisions:
        p = doc.add_paragraph(style='List Bullet')
        r = p.add_run(f'{label}: ')
        r.bold = True
        p.add_run(rule)
    doc.add_heading('Version update', level=2)
    doc.add_paragraph('v1.6 - Locked the My Programs layout and behavior, including ordering, card interaction, duplicate behavior, library ownership, large-collection controls, and the empty state.')

doc.save(SPEC)
