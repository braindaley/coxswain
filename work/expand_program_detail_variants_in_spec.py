from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches


SPEC = Path('/Users/brian/Documents/ChatGPT/Coxswain/outputs/Coxswain_Product_Design_Spec.docx')
ATLAS = Path('/Users/brian/Documents/ChatGPT/Coxswain/work/screen-atlas')

doc = Document(SPEC)
replacements = {
    'Primary actions: Create Program; View Program; Start; Edit; Duplicate; Export; Delete.':
        'Primary actions: Create Program; View Program; Start; Duplicate; Delete.',
    'Purpose: Create or edit a reusable program.':
        'Purpose: Create a reusable program or customize a duplicated definition.',
    'States and rules: Supports duration, distance, and interval definitions. Validate mixed-unit and empty-segment cases.':
        'States and rules: Supports duration, distance, and interval definitions. A saved program becomes immutable when created; changing it requires Duplicate so its existing workout History remains comparable. Validate mixed-unit and empty-segment cases.',
    'Primary actions: Start Workout or Start Race; view program History; open owned-program actions; back. Library programs use the same review structure without owned-program management actions.':
        'Primary actions: Start Workout or Start Race; view program History; Duplicate; Delete; back. Library programs use the same review structure without owned-program management actions.',
    'States and rules: Duration, distance, and interval programs reuse the same sections with type-specific values. Do not present an exact duration when distance-based work makes it unknowable. Race Your Best appears only with compatible completed workouts and replaces the ordinary start action while enabled.':
        'States and rules: Saved definitions are immutable so their History and best result always describe the same workout. Duplicate creates a new editable definition with no History. Delete removes the definition but retains completed workouts in History. Best is type-specific: timed programs use greatest distance; distance programs use fastest completion time; fixed-distance intervals use best average split across work segments; fixed-duration intervals use greatest work distance; mixed intervals use best average split across the identical ordered work segments. Rest is excluded from performance ranking. Race Your Best appears only with compatible completed workouts.',
    'Owned programs include View, Start, and an overflow management menu.':
        'Owned programs include View, Start, and an overflow menu containing only Duplicate and Delete.',
    '16.20 Program Detail / Start': '16.20 Program Detail — Duration',
    'Current v2.0 state: Program Detail / Start': 'Current v2.0 state: Program Detail — Duration',
    'The program definition is read-only and adapts to duration, distance, or interval content.':
        'Duration shows its fixed time target and completion rule without making the saved definition editable.',
    'Compatible program History enables Race Your Best; the primary action changes from Start Workout to Start Race.':
        'History ranks a timed program by greatest distance at the fixed elapsed time.',
    'Figure 24. Complete Coxswain screen inventory: 20 current screens and interaction states, plus 13 planned placeholders.':
        'Figure 26. Complete Coxswain screen inventory: 22 current screens and interaction states, plus 13 planned placeholders.',
}

found = set()
for paragraph in doc.paragraphs:
    source = paragraph.text.strip()
    if source in replacements:
        paragraph.text = replacements[source]
        found.add(source)
missing = set(replacements) - found
if missing:
    raise RuntimeError('Expected text was not found: ' + repr(sorted(missing)))

anchor = next(p for p in doc.paragraphs if p.text.strip() == '17 v2.0 Consolidated Decisions')

def add_variant(number, title, filename, bullets):
    heading_text = f'16.{number} Program Detail — {title}'
    if any(p.text.strip() == heading_text for p in doc.paragraphs):
        return
    heading = anchor.insert_paragraph_before(heading_text, style='Heading 2')
    heading.paragraph_format.page_break_before = True
    picture = anchor.insert_paragraph_before('', style='Normal')
    picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture.add_run().add_picture(str(ATLAS / filename), width=Inches(3))
    caption = anchor.insert_paragraph_before('', style='Normal')
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = caption.add_run(f'Current v2.0 state: Program Detail — {title}')
    run.italic = True
    anchor.insert_paragraph_before('Interaction changes', style='Heading 3')
    for bullet in bullets:
        anchor.insert_paragraph_before(bullet, style='List Bullet')
    anchor.insert_paragraph_before('', style='Normal')

add_variant(21, 'Distance', '21_program_detail_distance.png', [
    'Distance may show an estimated range while History ranks completions by fastest time.',
    'A compatible previous result enables Race Your Best.',
])
add_variant(22, 'Intervals', '22_program_detail_intervals.png', [
    'Intervals expose rounds, work, rest, total work, and segment guidance without claiming an exact duration.',
    'History ranks work-segment performance; Race Your Best requires the same ordered segment structure.',
])

doc.save(SPEC)
print('expanded Program Detail to duration, distance, and interval variants')
