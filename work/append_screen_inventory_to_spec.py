from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


SPEC = Path("/Users/brian/Documents/ChatGPT/Coxswain/outputs/Coxswain_Product_Design_Spec.docx")
BOARD = Path("/Users/brian/Documents/ChatGPT/Coxswain/work/screen-inventory-overview.png")

doc = Document(SPEC)
heading_text = "18. Complete Screen Inventory"
if any(paragraph.text.strip() == heading_text for paragraph in doc.paragraphs):
    raise RuntimeError("The complete screen inventory section already exists")

section = doc.add_section(WD_SECTION.NEW_PAGE)
section.orientation = WD_ORIENT.LANDSCAPE
section.page_width, section.page_height = section.page_height, section.page_width
section.top_margin = Inches(0.42)
section.bottom_margin = Inches(0.42)
section.left_margin = Inches(0.42)
section.right_margin = Inches(0.42)

heading = doc.add_paragraph()
heading.style = doc.styles["Heading 1"]
heading.add_run(heading_text)

intro = doc.add_paragraph(
    "This single-page inventory is the navigation map for the product design. Current tiles correspond to individual screenshots in Section 16; placeholder tiles reserve every screen that still needs design work. Replace a placeholder with its approved screenshot as the design is completed."
)
intro.paragraph_format.space_after = Pt(6)

image_paragraph = doc.add_paragraph()
image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
image_paragraph.paragraph_format.space_after = Pt(3)
image_paragraph.add_run().add_picture(str(BOARD), width=Inches(10.55))

caption = doc.add_paragraph(
    "Figure 23. Complete Coxswain screen inventory: 19 current screens and interaction states, plus 14 planned placeholders."
)
caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
for run in caption.runs:
    run.font.name = "Arial"
    run.font.size = Pt(8)
    run.font.italic = True
    run.font.color.rgb = RGBColor(83, 100, 124)

doc.save(SPEC)
print("appended the complete screen inventory to the specification")
