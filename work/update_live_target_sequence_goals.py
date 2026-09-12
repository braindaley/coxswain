from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches


SPEC = Path('/Users/brian/Documents/ChatGPT/Coxswain/outputs/Coxswain_Product_Design_Spec.docx')
ATLAS = Path('/Users/brian/Documents/ChatGPT/Coxswain/work/screen-atlas')
doc = Document(SPEC)

replacements = {
    'Recommended next screens to define and lock: 1) Program Builder, 2) Workout Library, 3) Program Detail / Start, 4) Live Row Target and goal states, 5) History and Workout Details.':
        'Recommended next screens to define and lock: 1) Race Reference Picker, 2) Pause / End, 3) Workout Complete, 4) History and Workout Details, 5) Connect Rower.',
    'Primary content: The approved configurable metric grid plus a compact progress panel. Duration shows elapsed and remaining time; distance shows completed and remaining meters; intervals show the current ordered segment, its progress, and the next segment.':
        'Primary content: The approved configurable metric grid plus a compact progress panel. Duration shows elapsed and remaining time; distance shows completed and remaining meters; intervals show the complete saved Row / Rest sequence on one line with the active segment highlighted.',
    'States and rules: The metric grid, Edit display interaction, Pause, and End session match Live Row Free. The progress panel is fixed below the grid. Free-form intervals advance through their saved order and never introduce rounds.':
        'States and rules: The metric grid, Edit display interaction, Pause, and End session match Live Row Free. The progress panel is fixed below the grid. Free-form intervals advance through their saved order and never introduce rounds. When Stroke rate, Speed, or Power has a goal, its metric card keeps the current reading as the large value and adds a compact variance bar with a labeled target inside the same card.',
    'Progress names the current Row or Rest segment, its remaining target, and the next ordered segment.':
        'The complete saved sequence appears on one line as Row, Rest, Row, Rest, and the active segment is highlighted.',
    'The screen follows the builder’s free-form segment order and has no rounds model.':
        'The active segment shows its remaining target. The screen follows the builder’s free-form order and has no rounds model.',
    'Figure 29. Complete Coxswain screen inventory: 25 current screens and interaction states, plus 12 planned placeholders.':
        'Figure 32. Complete Coxswain screen inventory: 28 current screens and interaction states, plus 12 planned placeholders.',
}

found = set()
for paragraph in doc.paragraphs:
    text = paragraph.text.strip()
    if text in replacements:
        paragraph.text = replacements[text]
        found.add(text)
missing = set(replacements) - found
if missing:
    raise RuntimeError('Expected text was not found: ' + repr(sorted(missing)))

anchor = next(p for p in doc.paragraphs if p.text.strip() == '17 v2.0 Consolidated Decisions')


def add_variant(number, title, filename, bullets):
    heading_text = f'16.{number} Live Row Target — {title}'
    if any(p.text.strip() == heading_text for p in doc.paragraphs):
        return
    heading = anchor.insert_paragraph_before(heading_text, style='Heading 2')
    heading.paragraph_format.page_break_before = True
    picture = anchor.insert_paragraph_before('', style='Normal')
    picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture.add_run().add_picture(str(ATLAS / filename), width=Inches(3))
    caption = anchor.insert_paragraph_before('', style='Normal')
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = caption.add_run(f'Current v2.0 state: Live Row Target — {title}')
    run.italic = True
    anchor.insert_paragraph_before('Interaction changes', style='Heading 3')
    for bullet in bullets:
        anchor.insert_paragraph_before(bullet, style='List Bullet')
    anchor.insert_paragraph_before('', style='Normal')


add_variant(26, 'Stroke-rate Goal', '26_live_target_goal_stroke_rate.png', [
    'The large 24 SPM value is the current stroke rate, not a difference from the goal.',
    'The bar places the current reading relative to the target zone and labels the 26 SPM target.',
])
add_variant(27, 'Speed Goal', '27_live_target_goal_speed.png', [
    'The large 13.8 km/h value is current speed.',
    'The same compact bar pattern places it relative to the labeled 14.5 km/h target.',
])
add_variant(28, 'Power Goal', '28_live_target_goal_power.png', [
    'The large 172 W value is current power.',
    'The same compact bar pattern places it relative to the labeled 180 W target.',
])

version_anchor = next(p for p in doc.paragraphs if p.text.strip() == 'Version update')
version_anchor.insert_paragraph_before(
    'Live performance goals: Stroke rate, Speed, and Power keep their current reading as the large card value. A compact in-card bar shows position relative to the labeled target; separate +/− variance metrics are not used.',
    style='List Bullet',
)

doc.save(SPEC)
print('updated interval sequence and added three live goal states')
