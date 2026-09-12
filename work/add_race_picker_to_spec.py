from pathlib import Path

from PIL import Image
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Inches, Pt, RGBColor


ROOT = Path('/Users/brian/Documents/ChatGPT/Coxswain')
SPEC = ROOT / 'outputs/Coxswain_Product_Design_Spec.docx'
SCREEN = ROOT / 'work/screen-atlas/30_race_reference_picker.png'
OVERVIEW = ROOT / 'work/screen-inventory-overview.png'

doc = Document(SPEC)

replacements = {
    'PLANNED': 'LOCKED',
    'Primary content: Compatible prior workouts, date, result, and comparison eligibility.':
        'Primary content: The current program and compatibility rule, selectable completed workouts, date, type-specific result, supporting pace, best-result marker, and an expandable list of excluded History.',
    'Primary actions: Select workout; review; start race; cancel.':
        'Primary actions: Select a compatible workout; inspect unavailable History; Start Race; Cancel or Back.',
    'States and rules: Filter to comparable workout structures and explain why an item is unavailable.':
        'States and rules: The best compatible result is selected by default. Distance programs compare identical target distance and rank fastest time. Duration programs compare identical target duration and rank greatest distance. Interval programs require the exact ordered Row and Rest structure and use the program-specific ranking rule. Unavailable History stays collapsed until requested and gives a concrete mismatch reason. Starting the race locks the selected result as the reference snapshot.',
    'Designed screens: Home, Quick Duration, Quick Distance, Quick Intervals, Segment Entry, Program Builder, My Programs, Workout Library, Program Detail / Start, Live Row Free, Live Row Target, and Live Row Race are represented in the atlas.':
        'Designed screens: Home, Quick Duration, Quick Distance, Quick Intervals, Segment Entry, Program Builder, My Programs, Workout Library, Program Detail / Start, Race Reference Picker, Live Row Free, Live Row Target, and Live Row Race are represented in the atlas.',
    'v2.0 - Consolidated all designed screens into an individual-state visual atlas, documented component changes by interaction, unified quick workout and program setup, simplified interval segments to Duration / Distance / Rest, and incorporated Race Your Best across setup and Live Row.':
        'v2.0 - Consolidated all designed screens into an individual-state visual atlas, documented component changes by interaction, unified quick workout and program setup, simplified interval segments to Duration / Distance / Rest, and incorporated Race Your Best across setup, reference selection, and Live Row.',
}

found = set()
inside_race_picker = False
for paragraph in doc.paragraphs:
    if paragraph.text == '10 Race Reference Picker':
        inside_race_picker = True
    elif paragraph.text == '11 Live Row Race Yourself':
        inside_race_picker = False
    original = paragraph.text
    if original == 'PLANNED' and not inside_race_picker:
        continue
    if original in replacements:
        paragraph.text = replacements[original]
        found.add(original)

missing = set(replacements) - found
if missing:
    raise RuntimeError('Expected specification text was not found: ' + repr(sorted(missing)))

# Refresh the all-screen inventory image while it is still the final inline image.
inventory = doc.inline_shapes[-1]
inventory_blip = inventory._inline.graphic.graphicData.pic.blipFill.blip
doc.part.related_parts[inventory_blip.embed]._blob = OVERVIEW.read_bytes()
pixel_width, pixel_height = Image.open(OVERVIEW).size
inventory.height = int(inventory.width * pixel_height / pixel_width)

anchor = next(p for p in doc.paragraphs if p.text == '17 v2.0 Consolidated Decisions')
inserted = []

page_break = doc.add_paragraph()
page_break.add_run().add_break(WD_BREAK.PAGE)
inserted.append(page_break)

heading = doc.add_paragraph('16.30 Race Reference Picker', style='Heading 2')
inserted.append(heading)

figure = doc.add_paragraph()
figure.alignment = WD_ALIGN_PARAGRAPH.CENTER
figure.add_run().add_picture(str(SCREEN), width=Inches(3.25))
inserted.append(figure)

caption = doc.add_paragraph('Current v2.0 state: Race Reference Picker')
caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
for run in caption.runs:
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(83, 100, 124)
inserted.append(caption)

subheading = doc.add_paragraph('Interaction changes', style='Heading 3')
inserted.append(subheading)

bullets = [
    'The screen shows only completed workouts compatible with the current program definition.',
    'The best compatible result is selected by default, while any listed result can become the fixed race reference.',
    'Result values adapt by program type: fastest time for distance, greatest distance for duration, and the correct work result for an identical interval sequence.',
    'Unavailable History is collapsed by default and explains each target or structure mismatch when opened.',
    'Start Race opens Live Row Race with the selected historical snapshot; Back and Cancel return without changing the reference.',
]
for text in bullets:
    inserted.append(doc.add_paragraph(text, style='List Bullet'))

for paragraph in inserted:
    anchor._p.addprevious(paragraph._p)

doc.save(SPEC)
print('added Race Reference Picker to the specification')
