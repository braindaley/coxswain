from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

SPEC = Path('/Users/brian/Documents/ChatGPT/Coxswain/outputs/Coxswain_Product_Design_Spec.docx')
HOME = Path('/Users/brian/Documents/ChatGPT/Coxswain/work/screen-atlas/01_home.png')
doc = Document(SPEC)

if doc.inline_shapes:
    shape = doc.inline_shapes[0]
    blip = shape._inline.graphic.graphicData.pic.blipFill.blip
    image_part = doc.part.related_parts[blip.embed]
    image_part._blob = HOME.read_bytes()

section = doc.sections[0]
header = section.header
hp = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
hp.text = 'COXSWAIN PRODUCT DESIGN SPEC  |  v2.0  |  CURRENT VISUAL ATLAS'
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
for run in hp.runs:
    run.font.name = 'Arial'
    run.font.size = Pt(8)
    run.font.bold = True
    run.font.color.rgb = RGBColor(83, 100, 124)

footer = section.footer
fp = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
fp.text = 'Living specification. Section 16 contains the current visual source of truth.'
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
for run in fp.runs:
    run.font.name = 'Arial'
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(83, 100, 124)

doc.save(SPEC)
print('cover image and section header/footer updated')
